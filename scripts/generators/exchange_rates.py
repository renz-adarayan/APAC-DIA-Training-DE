"""Exchange rates data generator module."""
import random
from datetime import timedelta, date
from pathlib import Path
from typing import Dict
from decimal import Decimal, ROUND_HALF_UP
import pyarrow as pa
from openpyxl import Workbook
from openpyxl.styles import NamedStyle, Font
from openpyxl.utils import get_column_letter

from utils.data_utils import apply_scale_to_targets, generate_date_range
from utils.schema_utils import get_column_names
from utils.constants import TARGET_ROWS


def generate_exchange_rates_data(schema: pa.Schema, scale: float, output_path: Path) -> int:
    """Generate exchange rates data and write to XLSX file.
    
    Args:
        schema: PyArrow schema for exchange_rates
        scale: Scaling factor for number of records
        output_path: Path to write the XLSX file
    """
    # Calculate target number of days (approximately 3 years)
    base_days = TARGET_ROWS['exchange_rates']  # 1,100 days
    num_days = apply_scale_to_targets(base_days, scale)
    
    # Ensure minimum meaningful data
    num_days = max(30, num_days)  # At least 1 month of data
    
    # Define currencies with realistic base rates to AUD
    currencies = {
        'USD': 0.6500,  # US Dollar
        'EUR': 0.5800,  # Euro
        'GBP': 0.5200,  # British Pound
        'JPY': 95.0000,  # Japanese Yen
        'NZD': 1.0800,  # New Zealand Dollar
        'CAD': 0.8800,  # Canadian Dollar
        'CNY': 4.7000,  # Chinese Yuan
        'SGD': 0.8700,  # Singapore Dollar
    }
    
    # Start date (3+ years ago) but cap end at 2024-12-31 to avoid 2025 spillover
    capped_end = date(2024, 12, 31)
    # If today is after capped_end, use capped_end; else use today (still <= capped_end)
    effective_end = min(date.today(), capped_end)
    start_date = effective_end - timedelta(days=num_days - 1)
    # Do not go earlier than 2022-01-01 just to keep dataset reasonable (optional clamp)
    earliest_allowed = date(2022, 1, 1)
    if start_date < earliest_allowed:
        start_date = earliest_allowed
    dates = generate_date_range(start_date, effective_end)
    
    # Track last weekday rates for weekend handling
    last_weekday_rates: Dict[str, Decimal] = {}
    
    # Create workbook and worksheet
    wb = Workbook()
    ws = wb.active
    ws.title = "Exchange Rates"
    
    # Add headers
    column_names = get_column_names(schema)
    for col_idx, header in enumerate(column_names, 1):
        ws.cell(row=1, column=col_idx, value=header)
    
    # Style headers
    header_style = NamedStyle(name="header")
    header_style.font = Font(bold=True)
    for col_idx in range(1, len(column_names) + 1):
        ws.cell(row=1, column=col_idx).style = header_style
    
    row_count = 0
    current_row = 2  # Start after headers
    
    for current_date in dates:
        is_weekend = current_date.weekday() >= 5  # Saturday=5, Sunday=6
        
        for currency, base_rate in currencies.items():
            if currency not in last_weekday_rates:
                # Initialize with base rate plus small random variation
                rate = Decimal(str(base_rate * (1 + random.uniform(-0.02, 0.02))))
                last_weekday_rates[currency] = rate
            
            if is_weekend:
                # Use last Friday's rate for weekends
                rate = last_weekday_rates[currency]
            else:
                # Generate new rate with daily volatility for weekdays
                prev_rate = last_weekday_rates[currency]
                
                # Apply daily change: typically ±0.5-2% for major currencies
                daily_change = random.uniform(-0.02, 0.02)
                
                # Occasional larger movements (market events)
                if random.random() < 0.05:  # 5% chance of larger move
                    daily_change = random.uniform(-0.05, 0.05)
                
                rate = prev_rate * Decimal(str(1 + daily_change))
                
                # Ensure rate stays positive and reasonable
                min_rate = Decimal(str(base_rate * 0.5))
                max_rate = Decimal(str(base_rate * 2.0))
                rate = max(min_rate, min(max_rate, rate))
                
                # Update last weekday rate
                last_weekday_rates[currency] = rate
            
            # Round to 8 decimal places as per schema
            rate = rate.quantize(Decimal('0.00000001'), rounding=ROUND_HALF_UP)
            
            # Write to Excel
            ws.cell(row=current_row, column=1, value=current_date)
            ws.cell(row=current_row, column=2, value=currency)
            ws.cell(row=current_row, column=3, value=float(rate))
            
            current_row += 1
            row_count += 1
    
    # Auto-adjust column widths
    for col_idx in range(1, len(column_names) + 1):
        column_letter = get_column_letter(col_idx)
        ws.column_dimensions[column_letter].auto_size = True
    
    # Ensure output directory exists
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    # Save workbook
    wb.save(output_path)
    
    return row_count
